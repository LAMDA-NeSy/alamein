from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "progress-assets"
OUTPUT = ROOT / "阿拉曼裁判工作室_阶段进展_2026-07-26.docx"

INK = "172824"
GREEN = "295B50"
BRASS = "A47D31"
SAND = "E9E0CD"
PALE = "F5F2E9"
MUTED = "65706C"
WHITE = "FFFFFF"
LINE = "C9C1B0"
RED = "8A3B34"

FONT = "Arial Unicode MS"
MONO = "Arial"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=11, color=INK, bold=False, italic=False, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic
    return run


def add_paragraph(doc, text="", *, size=11, color=INK, bold=False, italic=False,
                  before=0, after=6, line=1.10, align=None, keep=False, font=FONT):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    if text:
        set_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic, font=font)
    return p


def add_rich_paragraph(doc, parts, *, before=0, after=6, line=1.10, align=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    for text, options in parts:
        set_run(p.add_run(text), **options)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), size={1: 16, 2: 13, 3: 11.5}[level], color=GREEN if level < 3 else INK, bold=True)
    return p


def add_bullet(doc, text, *, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        set_run(p.add_run(bold_prefix), size=11, color=INK, bold=True)
        set_run(p.add_run(text[len(bold_prefix):]), size=11, color=INK)
    else:
        set_run(p.add_run(text), size=11, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    set_run(p.add_run(text), size=11, color=INK)
    return p


def add_callout(doc, label, text, fill=PALE, accent=BRASS):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent_dxa=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), accent)
    borders.append(left)
    tc_pr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    set_run(p.add_run(f"{label}  "), size=10, color=accent, bold=True)
    set_run(p.add_run(text), size=11.5, color=INK, bold=True)
    add_paragraph(doc, "", after=5)


def add_caption(doc, text):
    p = add_paragraph(doc, text, size=9, color=MUTED, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = False
    return p


def add_picture(doc, path, width=6.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    set_run(run, size=9, color=MUTED, font=MONO)


def build_architecture_diagram(path):
    width, height = 1800, 940
    image = Image.new("RGB", (width, height), "#F6F2E8")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    font_bold = ImageFont.truetype(font_path, 46, index=0)
    font_body = ImageFont.truetype(font_path, 30, index=0)
    font_small = ImageFont.truetype(font_path, 25, index=0)

    draw.text((90, 55), "一次 AI 行动如何进入游戏", font=font_bold, fill="#172824")
    draw.text((90, 115), "决策、规划、裁判与表现层彼此独立", font=font_body, fill="#65706C")

    boxes = [
        ("指挥层", "玩家 / 简单规则 / 复杂规则 / 外接模型", "#21483F"),
        ("行动意图", "移动目标 · 联合攻击 · 结束阶段", "#3A6B5F"),
        ("规划层", "自动寻路 · 候选动作 · 只读查询工具", "#9B762E"),
        ("规则裁判", "移动力 · ZOC · 堆叠 · CRT · 补给 · VP", "#7A3B34"),
        ("结果表现", "状态更新 · 动画 · 日志 · 存档 · 战报", "#304F5C"),
    ]
    x, y, box_w, box_h, gap = 95, 220, 1610, 105, 40
    for index, (title, detail, color) in enumerate(boxes):
        y0 = y + index * (box_h + gap)
        draw.rounded_rectangle((x, y0, x + box_w, y0 + box_h), radius=18, fill=color)
        draw.text((x + 38, y0 + 24), title, font=font_body, fill="#FFFFFF")
        draw.text((x + 330, y0 + 29), detail, font=font_small, fill="#F9F7F0")
        if index < len(boxes) - 1:
            cx = x + box_w // 2
            draw.line((cx, y0 + box_h, cx, y0 + box_h + gap - 8), fill="#A47D31", width=9)
            draw.polygon([(cx - 14, y0 + box_h + gap - 22), (cx + 14, y0 + box_h + gap - 22), (cx, y0 + box_h + gap)], fill="#A47D31")

    image.save(path, quality=95)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, GREEN, 16, 8),
        "Heading 2": (13, GREEN, 12, 6),
        "Heading 3": (11.5, INK, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.10

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(hp.add_run("ALAMEIN JUDGE STUDIO  |  PROJECT UPDATE"), size=8.5, color=MUTED, bold=True, font=MONO)

    footer = section.footer
    fp = footer.paragraphs[0]
    set_run(fp.add_run("2026.07.26  ·  v2026.07.12.95    "), size=8.5, color=MUTED, font=MONO)
    add_page_number(fp)


def build_document():
    diagram_path = ASSETS / "ai-judge-flow.png"
    build_architecture_diagram(diagram_path)

    doc = Document()
    style_document(doc)

    # Page 1: report cover
    p = add_paragraph(doc, "阶段进展", size=10, color=BRASS, bold=True, after=8, font=MONO)
    p.paragraph_format.keep_with_next = True
    add_paragraph(doc, "从规则数字化到可玩的 AI 兵棋系统", size=27, color=INK, bold=True, after=5, line=1.02, keep=True)
    add_paragraph(doc, "El Alamein Judge Studio", size=16, color=GREEN, bold=True, after=12, font=MONO)
    add_rich_paragraph(doc, [
        ("当前版本  ", {"size": 9.5, "color": MUTED, "bold": True, "font": FONT}),
        ("v2026.07.12.95", {"size": 9.5, "color": INK, "bold": True, "font": MONO}),
        ("    进展日期  ", {"size": 9.5, "color": MUTED, "bold": True, "font": FONT}),
        ("2026 年 7 月 26 日", {"size": 9.5, "color": INK, "bold": True, "font": FONT}),
    ], after=13)
    add_callout(
        doc,
        "一句话进展",
        "项目已从规则验证原型推进为可完成场景选择、移动、战斗、补给、胜负、存档和 AI 对局的网页兵棋系统。当前重点已转向对局可读性、AI 策略质量与战后复盘。",
    )
    add_picture(doc, ASSETS / "setup-v95.png", width=6.35)
    add_caption(doc, "当前首页：三个历史场景、双方阵营、存档与战役归档集中在同一入口")

    # Page 2: purpose and scope
    doc.add_page_break()
    add_heading(doc, "1. 我们正在做什么", 1)
    add_paragraph(
        doc,
        "El Alamein Judge Studio 以《El Alamein》原版规则为基础，目标不是做一个自由摆放棋子的地图编辑器，而是把传统桌面兵棋转化为可验证、可操作、可由 AI 参与的数字对局环境。",
    )
    add_paragraph(
        doc,
        "玩家或 AI 只负责提出行动：移动到哪里、哪些单位参加攻击、是否结束阶段。真正决定动作是否合法、移动消耗多少、战斗结果如何执行的，是同一套内置裁判。这样，人类玩家、规则 AI 和外接大模型始终在完全一致的规则环境中对局。",
        after=10,
    )
    add_heading(doc, "三个历史场景", 2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["场景", "先手", "回合", "主要差异"]
    for idx, text in enumerate(headers):
        set_cell_shading(table.rows[0].cells[idx], GREEN)
        p = table.rows[0].cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(text), size=9.5, color=WHITE, bold=True)
    rows = [
        ("七月 · 第一次阿拉曼", "轴心国", "7", "东进压力、阿拉曼方向与补给线"),
        ("九月 · 阿拉姆哈勒法", "轴心国", "7", "首回合特殊顺序与轴心国攻击加倍"),
        ("十月 · 第二次阿拉曼", "盟军", "15", "后期轴心国向西撤出与 VP 结算"),
    ]
    for row_idx, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        if row_idx % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, PALE)
        for idx, text in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(text), size=9.5, color=INK, bold=idx == 0)
    set_table_geometry(table, [2300, 1200, 800, 5060], indent_dxa=120)
    set_repeat_table_header(table.rows[0])
    add_caption(doc, "三个标准场景共享同一裁判内核，但保留各自回合顺序、特殊规则和胜利条件。")

    add_heading(doc, "已经形成的完整闭环", 2)
    for text in [
        "选择场景并为双方分别指定玩家、简单规则 AI、复杂规则 AI 或模型指挥官 AI。",
        "创建自动存档，按阶段完成移动、战斗、机械化移动、补给移动与回合结算。",
        "离开游戏时保存当前状态；终局后归档战役并保留最终胜负与对局日志。",
        "规则书与棋子图册独立成页，可查询专业术语、战斗结果表和单位作用。",
    ]:
        add_number(doc, text)

    # Page 3: player experience
    doc.add_page_break()
    add_heading(doc, "2. 从工具原型到游戏界面", 1)
    add_paragraph(
        doc,
        "前端已经从早期的规则工具面板改造成“北非战地指挥部”风格。地图保持为视觉主体，顶部集中回合、阶段与 VP，右侧操作台只显示当前阶段真正需要的功能。",
        after=8,
    )
    add_picture(doc, ASSETS / "game-v95.png", width=6.5)
    add_caption(doc, "当前对局界面：地图、选中单位、阶段轨道和移动操作在一个视图中完成")
    add_heading(doc, "玩家能直接感知到的变化", 2)
    for text in [
        "阶段更清楚：移动阶段不再堆放战斗控件，战斗阶段也不再显示无关移动操作。",
        "单位更易识别：双方使用不同外框；当前行动单位、不可行动友军和敌军具有不同状态。",
        "定位更直接：从“局面”点击单位，地图会将该棋子置于地图可视区域中央并高亮。",
        "失败可解释：无法移动时指出具体阻断步骤，并列出地形、堆叠等 MP 成本。",
        "战斗可跟随：联合进攻、骰点、赔率、CRT 结果和逐单位撤退都按步骤展示。",
    ]:
        add_bullet(doc, text)

    # Page 4: AI and judge
    doc.add_page_break()
    add_heading(doc, "3. AI 如何在规则内玩游戏", 1)
    add_picture(doc, diagram_path, width=6.45)
    add_caption(doc, "AI 不直接修改棋盘：所有动作最终都要回到同一裁判内核")
    add_heading(doc, "三类 AI", 2)
    add_rich_paragraph(doc, [
        ("简单规则 AI。", {"size": 11, "color": INK, "bold": True, "font": FONT}),
        ("从当前合法动作中快速选择评分最高的动作，适合作为基础对手和回归测试对象。", {"size": 11, "color": INK, "font": FONT}),
    ])
    add_rich_paragraph(doc, [
        ("复杂规则 AI。", {"size": 11, "color": INK, "bold": True, "font": FONT}),
        ("同时考虑目标距离、补给、控制区、地形、兵力保存、战斗赔率和胜利点价值。", {"size": 11, "color": INK, "font": FONT}),
    ])
    add_rich_paragraph(doc, [
        ("模型指挥官 AI。", {"size": 11, "color": INK, "bold": True, "font": FONT}),
        ("通过兼容 Chat Completions 的外部模型 API 接收局面，调用只读工具检查路径、战斗、补给和动作评分，再提交最终意图。", {"size": 11, "color": INK, "font": FONT}),
    ], after=9)
    add_heading(doc, "外接模型的一次决策", 2)
    for text in [
        "前端整理回合、阶段、胜利条件、双方单位、地图情报、近期日志和候选动作。",
        "模型可查询单位、格子、补给、路径和战斗结果，但所有工具只读。",
        "移动时模型优先只提交“单位 + 目标格”，系统负责寻找最优合法路线。",
        "最终动作由前端裁判二次验证；只有合法动作才写入状态并播放。",
    ]:
        add_number(doc, text)
    add_callout(
        doc,
        "战斗表现",
        "AI 能组织多个相邻单位联合攻击，也会在赔率明显不利时保存兵力。新版会明确显示参战单位、攻防值、赔率、随机骰点与战斗结果；若跳过，也会说明是没有合法目标还是预期战果不利。",
        fill="EEF3EF",
        accent=GREEN,
    )

    # Page 5: recent improvements and validation
    doc.add_page_break()
    add_heading(doc, "4. 最近一轮重点优化", 1)
    add_paragraph(doc, "本轮工作围绕一个目标展开：让玩家始终知道现在能做什么、刚才发生了什么、为什么会得到这个结果。", after=8)
    for prefix, text in [
        ("移动解释：", "显示准确阻断位置和 MP 成本明细。"),
        ("联合进攻：", "可依次选择多个己方单位，再点击敌军目标发起攻击。"),
        ("撤退选择：", "一次为一个单位选择完整路线，明确当前单位、格数与合法方向。"),
        ("AI 撤退：", "AI 控制的撤退方自动规划；人类控制方仍可选择合法路线。"),
        ("性能优化：", "可达范围与自动寻路转移到 Web Worker，并缓存高成本结果。"),
        ("信息降噪：", "删除重复面板和开发型入口，把规则、日志和调试放到更合适的位置。"),
    ]:
        add_bullet(doc, prefix + text, bold_prefix=prefix)

    add_heading(doc, "验证状态", 2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, text in enumerate(("检查项", "当前结果", "说明")):
        set_cell_shading(table.rows[0].cells[idx], INK)
        p = table.rows[0].cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(text), size=9.5, color=WHITE, bold=True)
    checks = [
        ("规则引擎自动测试", "29 项通过", "覆盖回合、移动、ZOC、堆叠、战斗、撤退、补给与 VP"),
        ("核心脚本", "语法检查通过", "前端、规则引擎和移动 Worker"),
        ("页面运行", "v95 正常", "控制台无错误，5176 端口可访问"),
        ("场景流程", "3 个场景", "七月/九月/十月的先手、回合和特殊顺序已接入"),
        ("AI 回放", "可到终局", "规则 AI 能执行移动、战斗和阶段推进"),
        ("外接模型协议", "链路已建立", "含上下文审计、质量检查、回放和最终动作复核"),
    ]
    for row_idx, values in enumerate(checks, start=1):
        cells = table.add_row().cells
        if row_idx % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, PALE)
        for idx, text in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(text), size=9.2, color=GREEN if idx == 1 else INK, bold=idx in (0, 1))
    set_table_geometry(table, [2400, 1900, 5060], indent_dxa=120)
    set_repeat_table_header(table.rows[0])
    add_caption(doc, "当前验证说明系统已从静态演示进入可持续迭代阶段。")

    # Page 6: next phase
    doc.add_page_break()
    add_heading(doc, "5. 仍在解决的问题", 1)
    for text in [
        "继续逐条核对原版规则的边缘情况，尤其是十月场景的撤出坐标与少数地图数据。",
        "扩大真人与模型对局样本，判断 AI 是否会形成有目的的战线，而不仅是完成合法动作。",
        "继续优化长对局中的 AI 节奏、战斗可读性和日志摘要。",
        "补充终局战报，包括关键战斗、VP 变化、损失和战线推进过程。",
        "完成移动端布局与低性能设备上的动画降级。",
        "在不暴露 API Key 的前提下完善本地代理、配置导入和部署说明。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "6. 下一阶段目标", 1)
    for text in [
        "建立固定的 AI 对局评测集，比较简单规则、复杂规则和外接模型在同一局面中的选择。",
        "为每次 AI 决策记录候选动作、选择理由、裁判结果和最终局面变化。",
        "强化战后报告，让一局结束后能够直接看到决定胜负的几次行动。",
        "完成三个场景的规则与数据复核，并形成可重复的发布检查流程。",
        "继续打磨首页、地图、战斗和终局之间的视觉一致性。",
    ]:
        add_number(doc, text)

    add_callout(
        doc,
        "下一阶段关键词",
        "从“规则正确、流程可用”推进到“对局有策略、结果可复盘”。",
        fill=SAND,
        accent=BRASS,
    )
    add_heading(doc, "结语", 1)
    add_paragraph(
        doc,
        "这段时间最大的变化，不只是界面变得更像游戏，而是系统已经形成清晰的职责边界：玩家和 AI 做决策，路径规划器寻找执行方式，裁判验证规则，前端负责把结果讲清楚。",
    )
    add_paragraph(
        doc,
        "这让项目具备了一个更有价值的下一步：我们不仅可以继续做一款可玩的阿拉曼兵棋，也可以把它作为研究 AI 如何在严格规则、部分信息和长期目标下进行决策的实验环境。",
        after=12,
    )
    add_paragraph(doc, "ALAMEIN JUDGE STUDIO", size=10, color=BRASS, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, font=MONO)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
